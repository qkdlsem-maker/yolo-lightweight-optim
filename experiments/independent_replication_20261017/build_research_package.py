"""Build the new empirical-study author-review package from verified JSON.
Original submitted/user-edited documents are never modified.
"""
from pathlib import Path
import json, hashlib, math, sys
sys.path.append(str(Path(__file__).resolve().parents[1]/'pydeps'))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

R=Path(__file__).resolve().parent
OUT=R.parents[1]/'outputs'/'Independent_Study_20261021'
OUT.mkdir(parents=True,exist_ok=True)
FIG=OUT/'figures'; FIG.mkdir(exist_ok=True)
def read(p): return json.loads(p.read_bytes())
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
A=read(R/'final_all_pair_analysis.json'); D=read(R/'completed_training_diagnostics_summary.json')
assert read(R/'final_analysis_independent_verification.json')['analysis_sha256']==sha(R/'final_all_pair_analysis.json')
G={x['group']:x for x in A['all_group_summaries']}
X={(x['arm'],x['weight'],x['epoch']):x for x in A['all_distillation_vs_control_contrasts']}
MET=['bdd_native_AP50_95','bdd_common_AP50_95','kitti_source_ignore_AP50_95','kitti_no_ignore_sensitivity_AP50_95','kitti_native_mean_moderate_AP_R40']
ARMS=[('mse',.03),('mse',.3),('mse',3),('cwd',.03),('cwd',.3),('cwd',3)]
labels=[f'{a.upper()} {w:g}' for a,w in ARMS]
TITLE='Feature distillation after detector fine tuning across three starting pairs and fixed evaluation definitions'
REPO='https://github.com/qkdlsem-maker/yolo-lightweight-optim'

plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False,'savefig.dpi':240})
fig,ax=plt.subplots(figsize=(7.2,3.6),layout='constrained')
colors=['#2166ac','#b35806','#5e3c99']
for i,seed in enumerate(['101','202','303']):
 vals=[X[(a,w,40)]['metrics'][MET[0]]['pair_values_pp'][seed] for a,w in ARMS]
 ax.scatter(vals,np.arange(6)+(i-1)*.13,s=32,c=colors[i],label=f'Pair {seed}',marker=['o','s','^'][i],zorder=4)
means=[X[(a,w,40)]['metrics'][MET[0]]['mean_pp'] for a,w in ARMS]
ax.scatter(means,np.arange(6),s=55,c='black',marker='|',label='Mean',zorder=5)
ax.axvline(0,c='#777777',lw=.8);ax.set_yticks(range(6),labels);ax.invert_yaxis()
ax.set_xlabel('Difference from the matched control (percentage points)')
ax.set_title('Primary BDD ten-class AP50:95 at epoch 40',loc='left',fontsize=12)
ax.grid(axis='x',color='#dddddd',lw=.6);ax.legend(ncol=4,loc='lower left',bbox_to_anchor=(0,1.06),frameon=False,fontsize=9)
fig.savefig(FIG/'Figure1_primary_pairs.png');fig.savefig(FIG/'Figure1_primary_pairs.pdf');plt.close(fig)

fig,axes=plt.subplots(1,2,figsize=(7.2,3.7),layout='constrained',sharey=True)
for ax,metric,title in zip(axes,[MET[2],MET[4]],['Common KITTI AP101\nsource ignore','Native KITTI AP_R40\nmean moderate']):
 for i,seed in enumerate(['101','202','303']):
  ax.scatter([X[(a,w,40)]['metrics'][metric]['pair_values_pp'][seed] for a,w in ARMS],np.arange(6)+(i-1)*.13,s=25,c=colors[i],marker=['o','s','^'][i])
 ax.axvline(0,c='#777777',lw=.8);ax.set_yticks(range(6),labels);ax.set_xlim(-.65,.65)
 ax.set_title(title,loc='left',fontsize=11);ax.set_xlabel('Difference (percentage points)');ax.grid(axis='x',color='#dddddd',lw=.6)
axes[0].invert_yaxis();fig.savefig(FIG/'Figure2_metric_definitions.png');fig.savefig(FIG/'Figure2_metric_definitions.pdf');plt.close(fig)

def font(run,name='Times New Roman',size=None,bold=False):
 run.font.name=name; run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
 if size:run.font.size=Pt(size)
 rpr=run._element.get_or_add_rPr(); rf=rpr.rFonts
 if rf is None:rf=OxmlElement('w:rFonts');rpr.insert(0,rf)
 rf.set(qn('w:eastAsia'),'Malgun Gothic')
 return run
def newdoc(short,korean=False):
 d=Document();s=d.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11)
 for st in d.styles:
  for elem in list(st.element.iter(qn('w:pBdr'))):elem.getparent().remove(elem)
 for elem in list(d.element.iter(qn('w:pBdr'))):elem.getparent().remove(elem)
 s.top_margin=s.bottom_margin=Inches(.72);s.left_margin=s.right_margin=Inches(.78)
 s.header_distance=s.footer_distance=Inches(.32)
 for nm in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
  st=d.styles[nm];st.font.name='Malgun Gothic' if korean else 'Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
  rf=st.element.get_or_add_rPr(); e=OxmlElement('w:rFonts');e.set(qn('w:eastAsia'),'Malgun Gothic');rf.append(e)
  st.paragraph_format.space_after=Pt(7)
 d.styles['Normal'].font.size=Pt(10.5 if korean else 11)
 d.styles['Normal'].paragraph_format.line_spacing=1.13
 d.styles['Title'].font.size=Pt(19);d.styles['Title'].font.bold=True
 for nm,sz in [('Heading 1',14),('Heading 2',12)]:
  d.styles[nm].font.size=Pt(sz);d.styles[nm].font.bold=True;d.styles[nm].paragraph_format.space_before=Pt(12)
 d.styles['Caption'].font.size=Pt(9.5);d.styles['Caption'].font.italic=False
 h=s.header.paragraphs[0];font(h.add_run(short),size=8)
 h.paragraph_format.space_after=Pt(0)
 f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 font(f.add_run('Author review draft  |  21 October 2026  |  '),size=8)
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
 d.core_properties.author='Hyerim Choi; Tae-Kook Kim';d.core_properties.title=TITLE if not korean else short
 return d
def p(d,t,style=None):return d.add_paragraph(t,style)
PAGE_PENDING=set()
def h(d,t,lev=1):
 pp=d.add_heading(t,lev)
 if id(d) in PAGE_PENDING:
  pp.paragraph_format.page_break_before=True;PAGE_PENDING.remove(id(d))
 return pp
def page(d):PAGE_PENDING.add(id(d))
def link(d,text,url, prefix=''):
 pp=d.add_paragraph(prefix);hlink=OxmlElement('w:hyperlink');hlink.set(qn('r:id'),pp.part.relate_to(url,RT.HYPERLINK,is_external=True))
 rr=OxmlElement('w:r');pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'20');pr.append(sz);rr.append(pr);tt=OxmlElement('w:t');tt.text=text;rr.append(tt);hlink.append(rr);pp._p.append(hlink);return pp
def table(d,headers,rows,widths=None,size=9.2):
 t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths is None:widths=[6.94/len(headers)]*len(headers)
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for i,x in enumerate(headers):t.rows[0].cells[i].text=str(x)
 for row in rows:
  for c,x in zip(t.add_row().cells,row):c.text=str(x)
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trpr.append(no)
  if ri==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for e in ['top','left','bottom','right']:
    x=OxmlElement('w:'+e);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
   pr.append(b);m=OxmlElement('w:tcMar')
   for e,val in [('top','65'),('bottom','65'),('left','75'),('right','75')]:
    x=OxmlElement('w:'+e);x.set(qn('w:w'),val);x.set(qn('w:type'),'dxa');m.append(x)
   pr.append(m)
   if ri==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E4EAF0');pr.append(shade)
   for para in c.paragraphs:
    para.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.05
    for rr in para.runs:font(rr,size=size,bold=ri==0)
 p(d,'').paragraph_format.space_after=Pt(0)
 return t
def caption(d,t):return p(d,t,'Caption')
def val(s,digits=3):return f"{s['mean_pp']:.{digits}f} ± {s['sample_SD_pp']:.{digits}f}"
def group(a,w,e):return G[f'{a}_w{w:g}_e{e}']
def row_name(r):
 if r['stage']=='base':return 'Base '+r['architecture']+' 100'
 return ('Control' if r['arm']=='control' else f"{r['arm'].upper()} {r['weight']:g}")+f"  e{r['epoch']}"

# Manuscript
d=newdoc('Feature distillation across starting pairs')
p(d,TITLE,'Title')
p(d,'Hyerim Choi¹ and Tae-Kook Kim²*')
p(d,'¹ Department of Artificial Intelligence Convergence, Pukyong National University, Busan 48513, Republic of Korea.\n² School of Computer and Artificial Intelligence Engineering, Pukyong National University, Busan 48513, Republic of Korea.')
p(d,'* Correspondence: king@pknu.ac.kr\nORCID for Hyerim Choi: 0009-0006-7311-4107')
h(d,'Abstract')
p(d,'Feature distillation can improve compact detectors, but its value after a substantial initial fine-tuning budget depends on the training intervention and evaluation definition. We studied a prespecified continuation regime using three independently fine-tuned YOLOv8 nano and large starting pairs. Each pair was trained for 100 epochs from fixed pretrained checkpoints. Seven matched student continuations then compared detection-only training with mean squared feature error and channel-wise distillation at coefficients 0.03, 0.3 and 3.0 for 40 epochs. All six base models and all 21 continuation paths at epochs 20 and 40 were evaluated, giving 48 fixed endpoints. We retained native ten-class BDD average precision, a common two-class evaluator on BDD and KITTI, and separate native KITTI scores. At epoch 40, the mean paired changes in BDD average precision were −0.0003, −0.0292 and −0.2597 percentage points for mean squared error, and +0.0096, −0.0497 and −0.2252 points for channel-wise distillation. Both coefficient-3 conditions decreased this metric in all three pairs. Common and native KITTI definitions yielded different descriptive effects; neither supports a causal explanation of transfer. Saved-state checks confirmed student updates and finite final states, while continuation mixed-precision counters recorded 916,981 applied and 299 skipped updates. The study provides an auditable sensitivity result for a narrow continuation regime, rather than a general ranking of distillation methods. Three pairs, reused validation populations, legacy annotation provenance and incomplete historical gradient instrumentation limit the conclusions.')
p(d,'Keywords: object detection; feature distillation; reproducibility; paired experiments; evaluation metrics')
page(d)
h(d,'Introduction')
p(d,'A compact detector is useful only if its accuracy and computational cost suit its intended setting. Knowledge distillation adds a training signal from a larger model without necessarily changing the deployed student. Its effectiveness is therefore often assessed as an accuracy difference against a student trained without that signal. Such a difference is interpretable only when starting weights, training budgets, data order and evaluation rules are sufficiently controlled. This is especially important when the student has already received substantial task-specific training and the remaining gains may be small.')
p(d,'Feature imitation is an established approach. Fine-grained imitation emphasizes detector responses near objects [1], while focal and global distillation separates local and relational information [2]. Channel-wise distillation (CWD) transfers spatial probability distributions within feature channels [3]. These methods address different aspects of the representation, and their published gains do not imply that every later adaptation or training schedule will improve. Cho and Hariharan examined limits of transferring knowledge from more accurate teachers [4]. Work on projectors also shows that alignment and normalization can affect optimization [5]. These studies motivate careful comparisons; they do not identify the cause of any result in the present experiment.')
p(d,'Recent work continues to change the distillation design. The September 2026 CSCWD preprint transfers high-resolution teacher features across scales in a YOLO11 aerial-detection setting [6]. Our experiment instead retains same-scale features in YOLOv8 and studies a fixed continuation budget. We do not reproduce CSCWD, claim to outperform it, or treat its reported improvements as directly comparable to the driving-image results below. Our contribution concerns the interpretation and reproducibility of a bounded experiment, not a new loss function or detector architecture.')
p(d,'Earlier project experiments varied continuation randomness while sharing a starting student–teacher pair; some still earlier implementation records were unsuitable for estimating distillation effects. Those observations motivated a new protocol with independently fine-tuned starting pairs and explicit loader seeds. They are not pooled into the present sample, and historical pruning comparisons are outside the present analysis. The new protocol was frozen before its six base and 21 continuation runs, and all scheduled outcomes are reported.')
p(d,'We ask how paired mean squared error (MSE) and CWD continuation effects vary across three starting pairs and a fixed coefficient grid, and how their description changes when evaluation definitions are aligned. The design separates a primary BDD ten-class endpoint from secondary common-class and native KITTI diagnostics. It also checks whether the intended intervention actually updated parameters and whether numerical summaries can be traced to retained artifacts. This combination supplies a focused empirical result for continued training after 100 epochs; it does not estimate a general effect across architectures, datasets or schedules.')
h(d,'Materials and methods')
h(d,'Data provenance and evaluation populations',2)
p(d,'BDD100K provides varied driving imagery and multiple visual tasks [7]. This study uses a reconstructed legacy 2018-format label mirror, audited against the pinned source JSON, rather than authenticated official-host Detection 2020 labels. Of 70,000 candidate training images, 137 with unresolved labels were excluded before the new experiments. The remaining 69,863 images contain 1,286,870 native training boxes after removal of one exact duplicate. The validation population has 10,000 images and 185,526 boxes. No unidentified local test labels are used. The original labels and earlier outputs remain unchanged.')
p(d,'BDD validation had already been used in earlier project work. KITTI evaluation uses all 7,481 images of its public labeled object-detection training partition, also examined previously in this project [8]. It is an external dataset relative to BDD training, but neither a newly untouched holdout nor the official hidden KITTI test set. No KITTI score was used to select a coefficient, checkpoint, epoch or replacement run in this fixed stage. Common-category BDD annotations contain 115,768 person/car targets. The KITTI common projection contains 33,229 targets and 25,726 explicit ignore records.')
h(d,'Independent starting pairs and paired continuations',2)
p(d,'Ultralytics 8.2.103 was pinned [9]. For pair seeds 101, 202 and 303, YOLOv8 nano and large models each received 100 epochs of BDD fine tuning from fixed official COCO-pretrained checkpoints. Independence refers to task-specific initialization and training randomness, not to independent COCO pretraining or independent dataset samples. The three matched starting pairs are the unit of replication. Initial parameter digests and sampled actual input tensors differed across seeds within each architecture.')
p(d,'Base training used 640-pixel input, SGD with initial learning rate 0.01, final multiplier 0.01, momentum 0.937 and weight decay 0.0005. Batch sizes were 64 for nano and 16 for large, with nominal batch size 64 and the native accumulation behavior. Warmup lasted three epochs, mosaic closed during the last ten epochs, early stopping was disabled, and mixed precision was enabled. The final epoch-100 EMA checkpoint was retained, never a validation-selected best checkpoint. A preflight found that the native training-loader generator did not follow the requested seed. Before full training, it was replaced with a finite DataLoader whose generator seed is the pair seed plus 100003 times the zero-based epoch; workers were recreated each epoch while native data transforms and optimizer behavior were retained.')
p(d,'Each nano endpoint initiated seven 40-epoch continuations: a detection-only control and MSE or CWD at coefficients 0.03, 0.3 and 3.0. The paired large endpoint served as the frozen evaluation-mode teacher. Continuation data seeds were the pair seed plus 1000; an epoch-dependent generator used that seed plus 100003 times the one-based epoch. All seven arms in a pair used the same image and augmentation sequence. The first two actual input batches of every epoch were hashed, including image tensors, classes, boxes and batch indices. All 63 within-pair arm comparisons agreed over the 40 epochs for these sampled batches; this is not an assertion that every batch was archived.')
p(d,'Continuation used SGD in one parameter group with momentum 0.937, weight decay 0.0005 and no Nesterov momentum. The initial learning rate was 0.0001, with a cosine schedule over 40 epochs and final multiplier 0.1. Batch size was 64, with four workers and a final batch of 39 images. Mosaic, mixup and copy-paste were disabled. Native HSV, translation, scale and horizontal-flip augmentation remained active. The student was trainable except for its fixed distribution-focal-loss projection. The control followed the same optimization schedule. Gradient clipping used norm 10; the mixed-precision scaler started at 128. EMA used decay 0.9999 and time constant 2000. Epoch 40 is primary, while epoch 20 is a secondary checkpoint on the same 40-epoch schedule, not a separately optimized schedule.')
h(d,'Distillation objectives and normalization',2)
p(d,'Features were tapped at layers 15, 18 and 21. Student channel counts were 64, 128 and 256; teacher counts were 256, 512 and 512. Fresh learned 1 × 1 convolutions aligned the student to teacher channels at each scale. MSE was the mean squared difference across all batch, channel and spatial entries, summed across the three scales. CWD applied a spatial softmax separately within each channel to aligned student and detached teacher features. It summed teacher-to-student Kullback–Leibler divergence over spatial positions, averaged over batch and channels, and multiplied by the squared temperature, with temperature fixed at 4. The three scale losses were summed [3].')
p(d,'The optimized objective was the native batch-summed detection loss plus the actual batch size multiplied by the fixed coefficient and the summed feature loss. This explicit scaling matters: equal coefficients do not imply equal MSE and CWD objective strengths, and loss magnitude does not measure gradient contribution. Adapters were optimized with the student and discarded for inference. Teacher computation was detached in the trainer; offline checks establish teacher-file integrity, not a reconstruction of past live gradients. Complete frozen configuration files and loss diagnostics accompany the code release.')
h(d,'Fixed evaluation definitions',2)
p(d,'The roster was fixed to six base checkpoints plus both epochs 20 and 40 of all 21 continuations, for 48 endpoints. Standalone inference used FP32, square 640 × 640 input, batch size 32, confidence threshold 0.001, NMS IoU 0.7 and at most 300 detections per image. Original class-aware NMS considered all ten classes before person/car mapping. Saved input-shape, precision and image-ledger records verified actual execution. Native training-context validation numbers were retained as diagnostics and did not replace standalone scores.')
p(d,'The primary metric is native Ultralytics ten-class BDD AP50:95. Native matching and integration were left unchanged. The common diagnostic uses pycocotools 2.0.11 COCOeval [10,11], the same continuous-coordinate box convention on both datasets, IoUs from 0.50 to 0.95 in steps of 0.05, and the mean interpolated precision at 101 recall points. The two classes are person and car, equally weighted; all areas and every image are retained. This common diagnostic is distinct from native BDD integration and from the official KITTI metric.')
p(d,'The primary common KITTI view maps Van and Person_sitting to class-specific ignore records and DontCare regions to both classes using COCO crowd intersection-over-detection-area semantics. A second, no-ignore sensitivity view omits those ignore records. Neither duplicates native KITTI difficulty and ignore rules. Separately, the pinned original KITTI 2D matching core applies Car IoU > 0.7 and Pedestrian IoU > 0.5 with its original difficulty, neighbor-class and DontCare behavior [8]. Native AP_R40 averages saved precision slots 1–40, excluding slot 0. We retain all easy, moderate and hard class scores and summarize native results by the mean of Car and Pedestrian moderate AP_R40.')
p(d,'Before new external scoring, known-answer tests covered common matching and integration, annotation projections, serialization, empty images, duplicates, ignored regions and interrupted archive recovery. Actual GPU input geometry and precision were checked on synthetic fixtures for both architectures. Synthetic results validate software behavior, not detector accuracy. Image and label hashes bind evaluation inputs to the frozen plan. Prediction archives preserve committed batches and reject mismatched sources or inputs, enabling technical recovery without repeating completed inference.')
h(d,'Analysis and verification',2)
p(d,'For each coefficient and objective, the paired effect is the student score minus its pair-matched control score at the same epoch. We report all three values, their arithmetic mean and sample standard deviation, in percentage points. The same three pairs contribute every coefficient and epoch; 21 trajectories and 48 endpoints are not 21 or 48 independent replications. No significance threshold, equivalence margin, best-method selection or historical pooling was added after seeing results. The prespecified six epoch-40 contrasts remain primary, and all six epoch-20 contrasts remain secondary.')
p(d,'Saved-state audits checked changed raw parameters, finite model/adapter/EMA/optimizer tensors and final checkpoint equality to saved EMA after conversion to checkpoint precision. Continuation scaler-step accounting recorded applied and skipped optimizer updates, and saved EMA counts agreed with applied updates. Numerical audits verified all 48 endpoint artifact chains, replayed native BDD saved statistics and the three common score variants, and recomputed native KITTI precision integration. These audits reuse pinned metric implementations; they are not independent implementations of the matching algorithms. A separate local arithmetic checker then compared 960 metric values and all 16 group summaries and 12 contrasts through 4,480 numerical comparisons at absolute tolerance 10⁻¹².')
h(d,'Results')
h(d,'Primary effects across all strengths',2)
p(d,'The mean native BDD AP50:95 was 25.8932 for the base nano models and 33.2661 for the base large models. The epoch-40 detection-only controls averaged 25.9488. Table 1 reports every primary contrast. MSE at 0.03 was close to zero on average, with different signs across pairs. CWD at 0.03 increased the metric in all three pairs, but by only 0.0022–0.0188 percentage points. Both objectives at coefficient 3 decreased native BDD AP in all pairs. These observations describe the tested runs and are not equivalence or population-significance conclusions.')
caption(d,'Table 1. All primary BDD ten-class AP50:95 contrasts at epoch 40. AP is reported on a 0–100 scale; differences and sample SD are percentage points. The control mean ± SD is 25.9488 ± 0.1666. Pair columns show treatment minus matched control.')
rows=[]
for a,w in ARMS:
 s=X[(a,w,40)]['metrics'][MET[0]]
 rows.append([f'{a.upper()} {w:g}',val(group(a,w,40)['metrics'][MET[0]],4),*[f"{s['pair_values_pp'][str(seed)]:+.4f}" for seed in [101,202,303]],f"{s['mean_pp']:+.4f} ± {s['sample_SD_pp']:.4f}"])
table(d,['Objective\ncoefficient','AP mean ± SD','Pair 101','Pair 202','Pair 303','Difference\nmean ± SD'],rows,[1.04,1.36,.96,.96,.96,1.66],8.8)
p(d,'Figure 1 shows the individual paired effects without treating standard deviation as a confidence interval. The largest mean primary reduction was 0.2597 percentage points for MSE at coefficient 3. All primary mean differences were within 0.26 points in magnitude. The small range should not be interpreted as a prespecified practical-equivalence finding; no such margin was defined.')
d.add_picture(str(FIG/'Figure1_primary_pairs.png'),width=Inches(6.85))
caption(d,'Figure 1. All three paired primary differences for each objective and coefficient. Colored points identify starting pairs; black ticks show means. Plotting code was prepared with OpenAI Codex assistance from the verified numeric table; no observations were generated or altered for the figure.')
h(d,'Common and native evaluation views',2)
p(d,'Table 2 retains the common BDD and both common KITTI variants alongside the native KITTI summary. At coefficient 3, common BDD differences were also negative on average: −0.1197 for MSE and −0.0890 for CWD. The corresponding common KITTI source-ignore changes were −0.0399 and +0.0099, whereas native KITTI mean-moderate changes were +0.0787 and +0.1061. These scores use different definitions and cannot be substituted for one another.')
caption(d,'Table 2. Epoch-40 paired differences under every secondary evaluation definition, mean ± sample SD in percentage points. Common AP uses two classes and 101 recall points. KI and KN respectively denote KITTI source-ignore and no-ignore sensitivity views. Native KITTI is mean moderate Car/Pedestrian AP_R40.')
table(d,['Objective\ncoefficient','BDD common','KITTI KI','KITTI KN','KITTI native'],[[f'{a.upper()} {w:g}',*[val(X[(a,w,40)]['metrics'][m],4) for m in MET[1:]]] for a,w in ARMS],[1.02,1.48,1.48,1.48,1.48],8.7)
p(d,'Figure 2 illustrates pair-level variability under the common source-ignore and native KITTI definitions. In particular, positive native mean differences can coexist with near-zero or negative common mean differences. The comparison demonstrates dependence on the chosen endpoint. It does not isolate annotation effects, class composition, scene distribution, teacher quality or domain shift as a cause. The source-ignore and no-ignore values are both retained instead of choosing the more favorable diagnostic.')
d.add_picture(str(FIG/'Figure2_metric_definitions.png'),width=Inches(6.85))
caption(d,'Figure 2. All epoch-40 paired KITTI effects under two distinct definitions. Colors and markers follow Figure 1. Both panels use the same axis range; individual points are three training pairs, not independent image samples. OpenAI Codex assisted with plotting code using the same verified observations.')
p(d,'The teacher bases averaged 42.1999 on common KITTI source-ignore AP and 76.7316 on native mean moderate AP_R40, compared with 38.9400 and 71.1691 for the student bases. Thus the larger models had higher averages on these definitions as well as native BDD. This descriptive teacher advantage does not establish how teacher representations affected the distillation gradients. Continued detection-only training raised the native BDD mean from 25.8932 to 25.9488 while lowering the native KITTI mean from 71.1691 to 70.6228; it is a contextual before/after observation, not a separately randomized comparison of training schedules.')
h(d,'Secondary endpoints and execution checks',2)
p(d,'All epoch-20 endpoints are retained in S1 Supplement and the machine-readable analysis. Their primary-metric paired mean differences were −0.0005, −0.0325 and −0.2774 for MSE, and −0.0076, −0.0082 and −0.2130 for CWD, in ascending coefficient order. The CWD 0.03 mean changed sign between epochs 20 and 40. Both high-coefficient means remained negative. We do not choose an epoch by its direction or interpret epoch 20 as a separate training replication.')
p(d,'All 27 scheduled training runs completed. Saved raw parameter tensors changed in every trainable student path, and final model, adapter, optimizer and EMA tensors were finite. Across the 21 continuations, 917,280 optimizer attempts comprised 916,981 applied updates and 299 mixed-precision skips. The first-batch weighted distillation-to-detection loss ratios averaged 0.00967, 0.09674 and 0.96735 for MSE, and 0.00721, 0.07215 and 0.72146 for CWD. These ratios explain the numerical scale of the tested objectives, not their causal gradient influence. All fixed-epoch loss summaries and every trajectory’s counters are retained in S1 Supplement and JSON.')
h(d,'Discussion')
p(d,'The principal result is conditional: adding these same-scale feature objectives to this conservative post-100-epoch schedule produced very small effects at coefficient 0.03 and negative mean BDD effects at larger coefficients. The negative coefficient-3 differences appeared in each starting pair for both losses. By changing task-specific starting pairs while preserving within-pair controls, the study goes beyond variation in continuation randomness around a single pair. However, three pairs remain a small sample, and all share pretrained sources and data populations.')
p(d,'This result does not contradict reports of successful detector distillation [1–3,6]. Those studies use their own architectures, loss designs and training regimes. Nor does it prove that feature distillation is ineffective after convergence: convergence itself was not established by a separate criterion, and the chosen continuation learning rate, adapters and augmentations form one specific intervention. We use “after fine tuning” to describe the observed 100-epoch starting point. Testing early-stage distillation, different projectors or other teachers would require separate, prospectively specified experiments. The present analysis offers no basis for selecting such settings using KITTI scores.')
p(d,'The two external score families answer different questions. Native KITTI applies class-specific overlaps and difficulty rules, whereas common AP averages ten IoU thresholds using a shared two-class implementation. Reporting both avoids presenting a favorable sign from one definition as universal transfer improvement. Even harmonized classes and integration do not align all annotation semantics or scene dependence. Our comparison therefore supports transparent metric reporting, not a causal mechanism for cross-dataset behavior.')
p(d,'The execution evidence also has a limited but concrete role. Parameter and state checks exclude the specific failure mode in which a completed student never changes. Actual scaler counts distinguish attempted from applied updates. Hashes and retained numerical statistics connect the published table to fixed files. These checks cannot validate the underlying research judgment, authenticate an unavailable official label release, or make an AI-assisted audit independent of its software assumptions. Reporting these boundaries is part of the result.')
h(d,'Limitations',2)
p(d,'The study includes one detector family, two feature objectives, three coefficients and three independently fine-tuned pairs. Shared COCO pretraining and reused BDD validation limit generalization; the legacy-label mirror is not authenticated Detection 2020 ground truth. Previously examined KITTI training images are not an untouched external test, and correlations among scenes are not modeled. No image bootstrap, confidence interval or significance/equivalence claim is made. The coefficient grid is a sensitivity analysis, not an optimization or method ranking.')
p(d,'Base training did not record actual per-step AMP applied/skipped counters, so they cannot be reconstructed from EMA totals. Only final raw continuation sidecars were retained; epoch-20 checkpoints are finite and hash-bound but lack the same independent full-precision EMA reconstruction as epoch 40. Per-batch gradient norms were not archived. There are 281 undefined or nonfinite epoch gradient-median records; their count does not reveal how many backward passes were nonfinite. Finite saved states and valid final updates do not prove that every intermediate gradient was finite. Numerical replay uses the pinned metric implementations, replays saved native BDD matches, and integrates native KITTI precision curves without independently repeating its C++ matching core.')
p(d,'Concurrent training wall time includes diagnostic evaluation and state writes. No isolated latency, energy, memory, pruning or deployment-safety benchmark was conducted in this stage. The students retain the nano inference architecture after adapter removal, but the study does not establish a new speedup. Full prediction archives and model binaries remain retained on the research server; the public package contains all endpoint metrics and hash-bound audit records, rather than a complete binary/raw-data archive. A durable access arrangement requires author review before journal submission.')
h(d,'Conclusions')
p(d,'Across three independently fine-tuned starting pairs, fixed MSE and CWD continuations yielded small low-coefficient effects and adverse high-coefficient BDD effects under the tested schedule. All 48 endpoints and all planned evaluation definitions are preserved. The contribution is a reproducible, bounded account of objective sensitivity and metric-dependent interpretation. Broader claims about distillation, teacher quality or domain shift require new evidence rather than a favorable subset of these results.')
h(d,'Data and code availability')
p(d,'The public project repository contains the frozen protocol, training and evaluation sources, synthetic validation records, all 48 endpoint scores, all paired contrasts, numerical audit proofs and figure source code. The new materials are labeled independent_replication_20261017 and are distinct from historical experiments. BDD and KITTI images and labels are obtained from their respective providers or the explicitly documented legacy mirror; they are not redistributed in this package. Checkpoint and prediction-archive hashes identify retained server artifacts. Full public deposition and permanent data access are not claimed in this author-review draft.')
link(d,'Public code and result repository',REPO)
h(d,'AI assistance')
p(d,'OpenAI Codex assisted with protocol and implementation development, code review, experiment orchestration, numerical analysis, literature and journal-policy retrieval, manuscript drafting and revision, and document and plotting code. Reported observations were computed by the recorded detector runs and evaluators. Automated checks included synthetic known-answer cases, artifact hashes, saved-state comparisons, numerical replay and a separate arithmetic calculation; these checks do not substitute for author validation. The figure legends identify assistance with plotting code. This draft records that assistance; the authors’ review and adoption of the analysis, references, disclosure and final wording remain pending.')
h(d,'Funding and author declarations')
p(d,'No research funding acknowledgment is included, following the author-provided funding information. The listed authors are Hyerim Choi and Tae-Kook Kim. Individual contribution roles, competing-interest statements and final author approval must be confirmed by both authors before submission; this draft does not assert that those declarations have been completed.')
h(d,'References')
refs=[
('Wang T, Yuan L, Zhang X, Feng J. Distilling Object Detectors with Fine-grained Feature Imitation. CVPR. 2019;4933–4942.','https://arxiv.org/abs/1906.03609'),
('Yang Z, et al. Focal and Global Knowledge Distillation for Detectors. CVPR. 2022;4643–4652.','https://openaccess.thecvf.com/content/CVPR2022/html/Yang_Focal_and_Global_Knowledge_Distillation_for_Detectors_CVPR_2022_paper.html'),
('Shu C, et al. Channel-wise Knowledge Distillation for Dense Prediction. ICCV. 2021;5311–5320.','https://arxiv.org/abs/2011.13256'),
('Cho JH, Hariharan B. On the Efficacy of Knowledge Distillation. ICCV. 2019;4794–4802.','https://arxiv.org/abs/1910.01348'),
('Miles R, Mikolajczyk K. Understanding the Role of the Projector in Knowledge Distillation. AAAI. 2024. arXiv:2303.11098v5.','https://arxiv.org/abs/2303.11098v5'),
('Zamani A, Ghasemi-Naraghi Z. CSCWD: Cross-Scale Channel-wise Knowledge Distillation for Lightweight Tiny Object Detection on Edge Devices. Preprint. 2026. arXiv:2609.30395v1.','https://arxiv.org/abs/2609.30395v1'),
('Yu F, et al. BDD100K: A Diverse Driving Dataset for Heterogeneous Multitask Learning. CVPR. 2020.','https://openaccess.thecvf.com/content_CVPR_2020/html/Yu_BDD100K_A_Diverse_Driving_Dataset_for_Heterogeneous_Multitask_Learning_CVPR_2020_paper.html'),
('KITTI. Object detection evaluation and original developer kit. Accessed 21 October 2026.','https://www.cvlibs.net/datasets/kitti/eval_object.php?obj_benchmark=2d'),
('Ultralytics. Release v8.2.103. Version pinned for this study.','https://github.com/ultralytics/ultralytics/releases/tag/v8.2.103'),
('Lin TY, et al. Microsoft COCO: Common Objects in Context. ECCV. 2014. arXiv:1405.0312.','https://arxiv.org/abs/1405.0312'),
('COCO API. Python COCOeval implementation. pycocotools 2.0.11 pinned in this study.','https://github.com/cocodataset/cocoapi/blob/master/PythonAPI/pycocotools/cocoeval.py'),
('BDD100K. Dataset download documentation distinguishing legacy and updated detection annotations.','https://github.com/bdd100k/bdd100k/blob/master/doc/source/download.rst'),
('Solesensei. BDD100K legacy data mirror used in the project annotation audit.','https://www.kaggle.com/datasets/solesensei/solesensei_bdd100k')]
for i,(t,u) in enumerate(refs,1):link(d,t,u,f'[{i}] ')
p(d,'Annotation provenance in Methods follows the project’s archived comparison to the legacy mirror [12,13]. The current review is a focused literature comparison, not a systematic review.')
h(d,'Supporting information')
p(d,'S1 Supplement. All 48 endpoint scores under five evaluation definitions; all native KITTI difficulty scores; both epochs’ paired summaries; objective-scale diagnostics; execution, recovery and verification boundaries. The accompanying JSON retains all 20 metrics per endpoint without rounding.')
d.save(OUT/'Manuscript_Independent_Distillation_20261021.docx')

# Supplement
s=newdoc('S1 Supplement  |  Independent distillation study')
p(s,'Supplementary results and verification records','Title')
p(s,'Hyerim Choi and Tae-Kook Kim  |  21 October 2026')
p(s,'This supplement accompanies the new study of three independently fine-tuned starting pairs. It contains every fixed endpoint, including unfavorable results. It does not pool earlier single-starting-pair experiments. Epoch 40 is primary; epoch 20 lies on the same 40-epoch schedule. All displayed AP values use a 0–100 scale, and differences use percentage points. Tables are rounded for readability; final_all_pair_analysis.json retains 960 unrounded metric values.')
h(s,'S1 Reading the endpoint tables')
table(s,['Column','Definition'],[['BDD native','Ten classes; unchanged Ultralytics AP50:95.'],['BDD common','Person/car; shared COCO AP101 at IoUs 0.50–0.95.'],['KITTI KI','Shared two-class AP101 with source-ignore projection.'],['KITTI KN','Same common metric without ignore records; sensitivity only.'],['KITTI native','Mean Car/Pedestrian moderate native AP_R40.']],[1.4,5.54],10)
p(s,'All-class inference precedes mapping. Shared classes and integration do not eliminate annotation, ignore or scene differences. KITTI is the previously examined 7,481-image public labeled training partition. BDD validation is reused and based on the reconstructed legacy mirror. Neither dataset is described as a fresh untouched test.')
h(s,'S2 Evidence chain')
p(s,'A frozen roster binds six base and 42 continuation checkpoints. Each completed endpoint binds all six execution steps; each numerical proof binds its completion marker, score artifacts and audit source. The completion manifests cover all 48 IDs exactly. Final numerical replay finished on 21 October 2026 at 00:05:48 UTC. A separate local arithmetic implementation passed at 00:12:40 UTC without importing the aggregation implementation or loading a model.')
for label,value in [('Protocol',sha(R/'protocol.json')),('Roster',A['roster_sha256']),('Plan',A['plan_sha256']),('Analysis',sha(R/'final_all_pair_analysis.json'))]:
 p(s,label+': '+value).paragraph_format.space_after=Pt(4)
p(s,'Software tests verify implementations under known inputs, and hash checks verify file identity. Neither establishes external validity or independent authentication of the label source. Raw copied training records, per-image data and private transport configuration are excluded from the public package.')
for seed in [101,202,303]:
 page(s);h(s,f'S3 All endpoints for pair {seed}')
 rows=[r for r in A['endpoint_rows'] if r['pair_seed']==seed]
 assert len(rows)==16
 caption(s,f'Table S{[101,202,303].index(seed)+1}. All 16 endpoints for starting pair {seed}. The first two rows are the epoch-100 bases; continuation labels give objective, coefficient and epoch. Scores are AP × 100.')
 table(s,['Endpoint','BDD\nnative','BDD\ncommon','KITTI\nKI','KITTI\nKN','KITTI\nnative'],[[row_name(r),*[f"{r['metrics_fraction'][m]*100:.4f}" for m in MET]] for r in rows],[1.94,1,1,1,1,1],9)
 p(s,'The same starting nano checkpoint is used by all seven arms within this pair. The corresponding large checkpoint is the teacher. Repeated rows across endpoints represent distinct fixed checkpoint times; they are not extra independent training pairs.')
 p(s,'The JSON companion additionally retains native BDD AP50, precision and recall, common per-class AP, native KITTI class/difficulty AP and every checkpoint SHA-256. No endpoint was selected by its score.')
for seed in [101,202,303]:
 page(s);h(s,f'S4 Native KITTI detail for pair {seed}')
 rows=[r for r in A['endpoint_rows'] if r['pair_seed']==seed]
 caption(s,f'Table S{[101,202,303].index(seed)+4}. Native AP_R40 × 100 for every endpoint. E, M and H denote easy, moderate and hard. This uses native class-specific overlaps and difficulty/ignore behavior, not common AP101.')
 nm=[f'kitti_native_{c}_{dif}_AP_R40' for c in ['car','pedestrian'] for dif in ['easy','moderate','hard']]
 table(s,['Endpoint','Car E','Car M','Car H','Ped E','Ped M','Ped H'],[[row_name(r),*[f"{r['metrics_fraction'][m]*100:.3f}" for m in nm]] for r in rows],[1.72,.87,.87,.87,.87,.87,.87],8.7)
 p(s,'The primary native KITTI summary is the arithmetic mean of Car M and Ped M; all remaining difficulty scores are preserved here. Detection outputs came from the same immutable KITTI prediction archive used for the common views. The numerical audit checked exact export/label bindings and recomputed AP_R40 from saved precision curves; it did not rerun native C++ matching.')
page(s);h(s,'S5 All secondary epoch contrasts')
caption(s,'Table S7. Epoch-20 paired differences, mean ± sample SD in percentage points. This is the same 40-epoch schedule, not a separate 20-epoch experiment. All primary epoch-40 contrasts are in manuscript Tables 1 and 2.')
table(s,['Objective\ncoefficient','BDD\nnative','BDD\ncommon','KITTI KI','KITTI KN','KITTI\nnative'],[[f'{a.upper()} {w:g}',*[val(X[(a,w,20)]['metrics'][m],3) for m in MET]] for a,w in ARMS],[1.04,1.18,1.18,1.18,1.18,1.18],8.3)
caption(s,'Table S8. Native BDD paired differences for epoch 20 in each starting pair, percentage points. Three pairs remain the sample size.')
table(s,['Objective\ncoefficient','Pair 101','Pair 202','Pair 303'],[[f'{a.upper()} {w:g}',*[f"{X[(a,w,20)]['metrics'][MET[0]]['pair_values_pp'][str(seed)]:+.4f}" for seed in [101,202,303]]] for a,w in ARMS],[1.84,1.7,1.7,1.7],10)
p(s,'For every other metric, the three raw paired differences are recoverable directly from Tables S1–S6 and are provided without rounding in all_distillation_vs_control_contrasts. Mean ± SD is descriptive across the three pairs. Neither a confidence interval nor a test of no effect is implied.')
page(s);h(s,'S6 Objective scale and training diagnostics')
p(s,'The first recorded batch supplies an exact weighted distillation-to-detection objective ratio. This is not a gradient ratio. Equal coefficients correspond to different loss scales. Saved first-batch snapshots were newly byte-bound during extraction; they were outside the earlier full saved-state audit and do not reconstruct historical live gradients.')
caption(s,'Table S9. First-batch weighted KD/detection loss ratio, mean ± sample SD across pairs. Also shown are epoch-40 detection and feature loss means. Detection is a mean of batch sums; feature loss is a mean of sums of per-scale means.')
rows=[]
for g in D['all_seven_group_summaries']:
 q=g['first_batch_weighted_KD_to_detection_objective_ratio'];e=next(x for x in g['fixed_epoch_loss_summaries'] if x['epoch']==40)
 ss=[e['mean_batch_summed_detection_loss'],e['mean_sum_of_feature_scale_means']]
 rows.append([g['group'].replace('_w',' '),f"{q['mean']:.5f} ± {q['sample_SD']:.5f}",*[f"{z['mean']:.4f} ± {z['sample_SD']:.4f}" for z in ss]])
table(s,['Objective\ncoefficient','First-batch ratio','Epoch 40\ndetection loss','Epoch 40\nfeature loss'],rows,[1.25,1.85,1.92,1.92],9)
p(s,'The last batch contains 39 images, while the other batches contain 64. The two epoch averages alone do not determine an exact epoch weighted-KD/detection ratio or an average of batch ratios. We therefore do not infer either. Epochs 1 and 20, all 840 epoch records summarized per trajectory, raw pair values and normalization notes are retained in completed_training_diagnostics_summary.json.')
p(s,'The 21 trainer wall times sum to 284,735.77 seconds, including diagnostic validation and state writes while other jobs ran concurrently. This is neither elapsed calendar duration nor an isolated GPU compute or inference benchmark. No energy or speedup claim is made from it.')
h(s,'S7 Applied and skipped updates')
p(s,'Across all continuations: 917,280 attempts = 916,981 applied updates + 299 skipped updates. The saved EMA update count matches the applied count for each trajectory. Skipped AMP steps are retained, not counted as successful parameter updates. Base runs did not retain equivalent per-step counters.')
page(s);h(s,'S8 Verification scope and retained failures')
p(s,'All six base final raw states, EMA states and optimizer tensors were audited. Within each architecture, starting and final parameter-value digests differed across the three seeds. Nano runs had 183 changed raw parameter tensors and large runs had 303; fixed DFL projections stayed unchanged. All 21 continuations showed changed student parameters and finite final model, adapters, EMA and optimizer states. Final checkpoints exactly matched the saved EMA after conversion to their serialization precision.')
p(s,'The final full-precision sidecars allow this comparison at epoch 40. Intermediate epoch-20 raw EMA sidecars were not retained. Those checkpoints are finite, architecture-checked and bound to trainer hashes, but the full-precision intermediate EMA cannot be independently reconstructed. For fixed DFL EMA values, tiny float32 accumulation differences were explained by a replay of the native EMA rule; they were not optimizer updates to frozen student parameters.')
p(s,'There are 281 undefined or nonfinite epoch gradient-median records. They remain explicit in the diagnostic output. The underlying per-batch norms were not archived, so this does not identify a frequency of nonfinite gradients, and it does not justify a claim that every backward pass was finite. Recorded finite epoch losses and final state checks have a different scope.')
h(s,'Technical recovery record',2)
p(s,'Before full training, a module-name collision and the fixed native loader generator were detected and corrected; original failed preflights remain in technical_failures. A base cache construction race was recovered under the unchanged protocol without repeating completed bases. Monitoring assumptions about exact EMA DFL equality, gradient medians and normal coordinator exit were corrected in separate audit records, preserving the original checks. None constitutes a discarded unfavorable experimental outcome.')
p(s,'During evaluation, the initial GPU command selected physical GPU 0 despite a masking environment variable. The two already-started native BDD jobs finished and their outputs were retained byte-for-byte. A separate scheduler version corrected physical-device selection for unfinished work. The cached large-pair-101 native BDD result is the documented GPU assignment exception. Scientific inference settings did not change, completed inference was not repeated, and no timing comparison is based on the affected jobs.')
p(s,'The first aggregation attempt used a folder name where a proof-field name differed. It failed before writing a result. A separate corrected source version mapped kitti_native to native_kitti and passed all input-hash and completeness checks. The original source and failure log remain in audit_technical_failures/004_aggregation_artifact_key. A second independent arithmetic implementation reproduced all final summaries. No scores, checkpoints or method settings were altered to resolve this reporting error.')
h(s,'Audit limits and reproducibility',2)
p(s,'Common scores were replayed with the same pinned COCO implementation. Native BDD numerical replay uses stored true-positive statistics instead of independently reconstructing IoU matches from images. Native KITTI replay recomputes integration from saved precision curves; it does not independently reproduce C++ matching. Source bindings and synthetic tests add evidence but do not create an independent implementation. Full raw datasets and binary prediction/model archives are not included in this public text/JSON package. Author review must settle durable access before submission.')
s.save(OUT/'Supplement_All_Endpoints_20261021.docx')

# Author-facing review, Korean.
v=newdoc('독립 시작 쌍 실험 검토와 제출 전 확인',True)
p(v,'독립 시작 쌍 실험 검토와 제출 전 확인','Title')
p(v,'2026년 10월 21일  |  Hyerim Choi · Tae-Kook Kim')
h(v,'현재 판단')
p(v,'새 단계의 기초 학습 6회, 후속 학습 21회와 고정된 48개 모델 시점의 평가·수치 감사가 완료되었다. 새 원고는 독립 시작 쌍과 고정 가중치의 민감도, 평가 정의에 따른 해석 차이를 다루는 제한된 실증 연구로 다시 작성했다. 새로운 증류 알고리즘, 일반적인 증류 실패, 교사 품질이나 도메인 차이의 인과효과를 주장하지 않는다. IEEE Access의 이전 원고는 사전 검토에서 최종 거절되었으며, 이 자료는 같은 원고의 재투고나 게재를 약속하지 않는다.')
h(v,'검증된 주요 결과')
table(v,['방법','가중치 0.03','가중치 0.3','가중치 3'],[['MSE','−0.0003 pp','−0.0292 pp','−0.2597 pp'],['CWD','+0.0096 pp','−0.0497 pp','−0.2252 pp']],[1.24,1.9,1.9,1.9],10)
p(v,'표는 40epoch BDD 10클래스 AP50:95의 대조군 대비 평균 차이다. CWD 0.03은 세 쌍 모두 양수였지만 차이가 0.0022–0.0188 pp로 매우 작았다. 가중치 3에서는 두 방법 모두 세 쌍의 점수가 낮았다. 표준편차는 기술통계이며 통계적 유의성이나 실질적 동등성 판정이 아니다. KITTI의 공통 지표와 원래 지표는 다른 결과를 낼 수 있어 모두 보존했다.')
p(v,'별도 집계 검산은 48개 감사 증거의 해시를 연결한 뒤 960개 원값, 16개 그룹, 12개 대비를 4,480회 수치 비교해 일치를 확인했다. 연속 학습의 실제 AMP 기록은 917,280회 시도 중 916,981회 적용, 299회 건너뜀이다. 기초 학습의 같은 카운터는 기록하지 않아 복원할 수 없다.')
h(v,'추가 실험 판단')
p(v,'이번 질문을 답하는 사전 명시 실험은 모두 끝났으므로, 유리한 결과를 얻기 위한 추가 학습은 실시하지 않았다. 새 구조, 초기 학습부터의 증류, 다른 projector 또는 인증된 새 라벨은 별도 연구 질문이다. 이를 조사하려면 조건·예산·종료 기준과 새 평가 계획을 먼저 정해야 한다. 현재의 작은 효과나 불리한 결과 자체는 실험을 늘릴 이유가 아니다. 기존 결과를 유지한 채 제한된 실증 기여로 제시하는 것이 현재 증거에 맞다.')
page(v);h(v,'선행연구 대비 기여')
p(v,'FGFI와 FGD는 검출 특징의 선택·관계 전이에 관한 방법 연구이고, CWD는 채널별 공간 분포 전이를 제안한 연구이다. Cho와 Hariharan은 큰 교사의 효과가 단순하지 않음을, Miles와 Mikolajczyk는 projector와 정규화의 역할을 조사했다. 이들은 본 실험의 실패 원인을 입증한 근거가 아니다. 2026년 CSCWD는 YOLO11의 교차 스케일 전이를 다루는 사전공개 논문으로, 이번 YOLOv8 동일 스케일 후속 학습과 직접 성능을 비교하지 않는다. 원고 참고문헌에 각각의 원문을 연결했다.')
p(v,'이번 기여는 3개의 독립적인 BDD 미세조정 시작 쌍, 6개의 고정 증류 조건, 쌍별 대조군, 전 모델 공통·원래 평가 지표, 실제 적용 업데이트 및 결과 파일의 추적 가능성을 함께 보고한 점이다. 새 방법의 성능 우위를 중심으로 평가하는 저널에는 기여가 약하다. 단순히 문장을 바꾸어 높은 신규성으로 표현하면 안 된다.')
h(v,'새 저널 적합성 검토')
p(v,'우선 검토 후보는 PLOS One이다. 출판 기준은 연구의 방법론적 타당성을 강조하며 부정적·재현 결과도 다룬다. 다만 기존 연구와의 차이와 질문의 가치를 설득해야 하므로, 이 범위의 실증 결과가 자동으로 충분하다는 뜻은 아니다. 현재 제공본은 저널 공통의 저자 검토용 원고이며 PLOS 제출 형식 충족을 선언하지 않는다. 최종 저널이 정해지면 제목 페이지·줄 번호·그림 별도 업로드·커버레터 등 제출 형식을 맞춰야 한다.')
link(v,'PLOS One 출판 기준','https://journals.plos.org/plosone/s/criteria-for-publication')
link(v,'PLOS One 제출 지침','https://journals.plos.org/plosone/s/submission-guidelines')
p(v,'Applied Sciences는 응용 실험 연구 범위의 대안이다. 출판사 검색 결과의 현재 색인 안내는 SCIE를 명시한다. PLOS One의 이번 직접 확인은 저널 범위·출판 기준·AI 정책까지이며, 최신 Clarivate SCIE 명단의 직접 조회를 완료했다고 주장하지 않는다. SCIE가 투고의 필수 조건이면 두 후보의 현재 Clarivate Master Journal List 상태를 투고 직전에 별도로 확인해야 한다. 게재 가능성이나 일정·비용을 보장하지 않는다.')
link(v,'Applied Sciences 범위','https://www.mdpi.com/journal/applsci/about')
link(v,'Applied Sciences 색인 안내','https://www.mdpi.com/journal/applsci/indexing')
p(v,'가장 큰 심사 위험은 3쌍이라는 작은 반복 수, 하나의 검출기 계열, 재사용 검증셋, legacy 라벨 출처, 매우 제한된 효과와 공개 원시 자료 접근성이다. 큰 교사의 평균 점수가 더 높다는 사실은 증류 작동 원인을 입증하지 않는다. 공통 지표를 사용해도 두 데이터셋의 라벨 의미와 장면 의존성이 같아지지 않는다. 이 한계들은 원고·보충자료에 모두 남겼다.')
page(v);h(v,'감사의 글과 AI 공개 상태')
p(v,'사용자가 직접 수정한 Desktop의 IEEE_Access_20261006.docx는 변경하지 않았다. 그 파일의 감사 문구는 다음과 같다.')
p(v,'“OpenAI Codex [8] assisted with code review, numerical analysis, and English editing and formatting of the Abstract and Sections I–VI. The reported results were computed from recorded YOLOv8 experimental data on BDD100K and KITTI. Source code, experiment records, and analytical outputs are provided to support verification. The authors are responsible for verifying the implementation, data, references, and final manuscript.”')
p(v,'이번 후속 작업에는 프로토콜·구현 개발, 실험 실행 지원, 분석·그림 코드, 문헌 검토와 새 원고 작성도 포함되어 과거 감사 문구만으로 이번 작업 전체가 설명되지는 않는다. 새 원고는 실제 지원 범위를 별도 AI assistance 절에 기술했다. 자동 검증 방법을 명시하되, 저자가 모든 내용·참고문헌·코드·해석을 검토하고 승인했다고 쓰지 않았다. 사용자가 과거 문구를 수정한 사실과 새 저널 요건이 충족되었는지는 서로 다른 문제이다.')
p(v,'PLOS의 현재 정책은 연구·내용 작성에 사용한 도구 이름, 영향을 받은 범위, 산출물의 타당성을 어떻게 확인했는지 본문에 공개하도록 한다. 코드 개발과 그림 준비도 해당한다. 최종 저자는 공개 문구를 실제 사용 내역과 대조하고 자신의 검증 내용을 확정해야 한다. 이 검토 기록은 저자의 승인이나 독립적인 사람의 동료심사를 대체하지 않는다.')
link(v,'PLOS AI 도구와 공개 정책','https://journals.plos.org/plosone/s/ethical-publishing-practice')
h(v,'제출 전 저자 확인 사항')
p(v,'두 저자가 연구 해석·인용문헌·저자별 기여·이해상충·최종 AI 공개를 직접 검토해야 한다. 저자 이름은 Hyerim Choi와 Tae-Kook Kim 두 명, 제공된 Hyerim Choi ORCID는 0009-0006-7311-4107로 유지했다. 연구비 사사는 넣지 않았다. 개인별 기여와 이해상충 여부는 새로 추정하지 않았다.')
p(v,'공개 패키지는 모든 모델 시점의 수치와 해시, 코드, 감사 결과를 제공한다. 전체 체크포인트·예측 바이너리와 데이터 이미지는 연구 서버에 보존되어 있으나 이 문서/JSON 공개물에는 포함되지 않는다. 데이터 제공자의 접근·사용 조건과 영구 보관 위치를 저자가 확인하고, 실제로 제공 가능한 범위에 맞게 Data availability 문구를 확정해야 한다. 비공개 접속 정보나 원시 이미지·라벨은 GitHub에 올리지 않는다.')
p(v,'어떤 저널에도 투고·메일 발송·비용 결제하지 않았다. 현재 파일은 검토와 수정이 가능한 새 연구 원고이며, 저자 승인 또는 게재 가능성을 보증하는 제출 완료본이 아니다. 기존 IEEE 원고와 이전 9회 실험의 결과는 그대로 보존하고 새 단계와 섞지 않았다.')
v.save(OUT/'Author_Review_and_Journal_Fit_20261021.docx')

manifest={'status':'authored_pending_full_page_visual_QA','analysis_sha256':sha(R/'final_all_pair_analysis.json'),
 'builder_sha256':sha(Path(__file__)),'authors':['Hyerim Choi','Tae-Kook Kim'],'orcid':'0009-0006-7311-4107',
 'funding_acknowledgment':False,'author_approval':'pending','AI_disclosure_author_verification':'pending',
 'files':{f.name:sha(f) for f in OUT.glob('*.docx')}}
(R/'authored_package_manifest.json').write_text(json.dumps(manifest,indent=2)+'\n',encoding='utf-8')
print(json.dumps(manifest,indent=2))
